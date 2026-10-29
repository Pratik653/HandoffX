"""Assemble and build the final consolidated HandoffX paper.

    python publication_final/build_final.py            # Markdown + DOCX + PDF
    python publication_final/build_final.py --md-only  # Markdown only

Sources: `src/*.md` (new text, in order) with include directives
`<<SRC|START|END[|DEMOTE][|keepstart][|sec=A:B]>>` that copy verbatim blocks from
the preserved manuscripts (SRC = v01: 01_Publication/HandoffX_Research_Paper.md,
v2: publication_v2/PART_II_v0.2.md). START/END are line prefixes; the block runs
from START (excluded unless `keepstart`) to the next line starting with END.
Tables and figures are numbered by order of definition ({T:key}, {F:key}).

Tooling only, not part of the handoffx package. Requires python-docx, latex2mathml,
lxml, Microsoft Office's MML2OMML.XSL and Microsoft Edge (as publication_v2/build_manuscript.py).
"""
from __future__ import annotations
from pathlib import Path
import html as H
import importlib.util, re, subprocess, sys

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
SOURCES = {'v01': ROOT / '01_Publication/HandoffX_Research_Paper.md', 'v2': ROOT / 'publication_v2/PART_II_v0.2.md'}
NAME = 'HandoffX_Final_Research_Paper'

# Wording updates applied to included historical text (each must match exactly once overall).
SUBSTITUTIONS = [
    ('The publication requires human authorship approval, licensing decisions and independent scientific/editorial review; it has not been deposited or assigned a DOI.',
     'Formal publication still requires confirmation of authorship metadata and independent scientific review; no DOI has been assigned at the time of writing.'),
]


# ------------------------------------------------------------------ assembly
def _strip_fm(text: str) -> str:
    return re.sub(r'^---\n.*?\n---\n', '', text, flags=re.S)


def _block(src: str, start: str, end: str, demote: int, keep: bool, sec: str | None) -> str:
    lines = _strip_fm(SOURCES[src].read_text(encoding='utf-8')).split('\n')
    i = next(k for k, l in enumerate(lines) if l.startswith(start))
    j = next(k for k in range(i + 1, len(lines)) if lines[k].startswith(end))
    out = lines[i if keep else i + 1:j]
    res = []
    for l in out:
        m = re.match(r'^(#{1,6}) (.*)$', l)
        if m:
            t = m.group(2)
            if sec:
                a, b = sec.split(':'); t = re.sub(rf'^{a}\.(\d+)\.', rf'{b}.\1.', t)
            pm = re.match(r'^P(\d)\. (.*)$', t)
            if pm and src == 'v01' and len(m.group(1)) == 2:
                t = f'7.{pm.group(1)}. {pm.group(2)} (P{pm.group(1)})'
            l = '#' * min(6, len(m.group(1)) + demote) + ' ' + t
        res.append(l)
    text = '\n'.join(res).strip('\n')
    if src == 'v01':   # v0.1 table/figure captions -> placeholders
        text = re.sub(r'\*\*Table (\d+)\*\*\n\n\*(.+?)\*', lambda m: f'**Table {{T:v1-{m.group(1)}}}.** *{m.group(2)}*', text)
        text = re.sub(r'\*\*Figure (\d+)\*\*\n\n\*(.+?)\*\n\n!\[\]\((.+?)\)', lambda m: f'![Figure {{F:v1-{m.group(1)}}}. {m.group(2)}]({m.group(3)})', text)
        text = re.sub(r'(?<![{\w])Figure (\d+)(?=[ .,;)])', r'Figure {F:v1-\1}', text)
        text = re.sub(r'(?<![{\w])Table (\d+)(?=[ .,;)])', r'Table {T:v1-\1}', text)
    return text


def assemble() -> str:
    parts = [p.read_text(encoding='utf-8') for p in sorted((HERE / 'src').glob('*.md'))]
    md = '\n\n'.join(parts)

    def inc(m):
        f = m.group(1).split('|')
        src, start, end = f[0], f[1], f[2]
        demote = int(f[3]) if len(f) > 3 and re.fullmatch(r'[+-]?\d+', f[3]) else 0
        return _block(src, start, end, demote, 'keepstart' in f, next((x[4:] for x in f if x.startswith('sec=')), None))
    md = re.sub(r'^<<(.+?)>>$', inc, md, flags=re.M)
    for old, new in SUBSTITUTIONS:
        n = md.count(old)
        if n != 1:
            raise SystemExit(f'substitution matched {n} times: {old[:60]}')
        md = md.replace(old, new)
    # number tables and figures by order of definition
    num = {}
    for kind, pat in (('T', r'\*\*Table \{T:([\w-]+)\}\.\*\*'), ('F', r'!\[Figure \{F:([\w-]+)\}\.')):
        for k, key in enumerate(dict.fromkeys(re.findall(pat, md)), 1):
            num[(kind, key)] = k
    def rep(m):
        key = (m.group(1), m.group(2))
        if key not in num:
            raise SystemExit(f'undefined reference {m.group(0)}')
        return str(num[key])
    md = re.sub(r'\{([TF]):([\w-]+)\}', rep, md)
    md = re.sub(r'\n{3,}', '\n\n', md)
    return md.strip() + '\n'


# ------------------------------------------------------------------ rendering helpers
def _load_v2_helpers():
    spec = importlib.util.spec_from_file_location('bm', ROOT / 'publication_v2/build_manuscript.py')
    bm = importlib.util.module_from_spec(spec); spec.loader.exec_module(bm)
    return bm


def _meta(md):
    fm, body = re.match(r'^---\n(.*?)\n---\n(.*)$', md, re.S).groups()
    return dict(re.findall(r'^(\w+):\s*"?(.*?)"?$', fm, re.M)), body


def _slug(t):
    return re.sub(r'[^a-z0-9]+', '-', t.lower()).strip('-')


BREAK_BEFORE = ('References', 'Appendix', 'Abstract', 'Declarations')


def build_docx(md: str, out: Path, bm):
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
    from docx.shared import Inches, Pt, RGBColor
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(); meta, body = _meta(md)
    sec = doc.sections[0]; sec.left_margin = sec.right_margin = Inches(1); sec.top_margin = sec.bottom_margin = Inches(1)
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(6); st.paragraph_format.line_spacing = 1.15
    for lvl, size in ((1, 15), (2, 12.5), (3, 11.5)):
        hs = doc.styles[f'Heading {lvl}']; hs.font.size = Pt(size); hs.font.color.rgb = RGBColor(0, 0, 0)
        rf = hs.element.get_or_add_rPr().get_or_add_rFonts()
        for a in ('w:asciiTheme', 'w:hAnsiTheme', 'w:eastAsiaTheme', 'w:cstheme'):
            rf.attrib.pop(qn(a), None)
        rf.set(qn('w:ascii'), 'Times New Roman'); rf.set(qn('w:hAnsi'), 'Times New Roman')

    def field(par, instr, placeholder=''):
        r = par.add_run(); b = OxmlElement('w:fldChar'); b.set(qn('w:fldCharType'), 'begin'); r._r.append(b)
        r = par.add_run(); t = OxmlElement('w:instrText'); t.set(qn('xml:space'), 'preserve'); t.text = instr; r._r.append(t)
        r = par.add_run(); s = OxmlElement('w:fldChar'); s.set(qn('w:fldCharType'), 'separate'); r._r.append(s)
        par.add_run(placeholder)
        r = par.add_run(); e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), 'end'); r._r.append(e)

    # footer page numbers; ask Word to refresh fields (TOC) on open
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER; field(fp, 'PAGE', '1')
    upd = OxmlElement('w:updateFields'); upd.set(qn('w:val'), 'true'); doc.settings.element.append(upd)
    # title page
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; r = t.add_run(meta.get('title', '')); r.bold = True; r.font.size = Pt(26)
    for k, size, it in (('subtitle', 13, True), ('author', 12, False), ('version', 10.5, False), ('date', 10.5, False)):
        if k in meta:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; rr = p.add_run(meta[k]); rr.italic = it; rr.font.size = Pt(size)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    c = doc.add_paragraph(); r = c.add_run('Contents'); r.bold = True; r.font.size = Pt(15)
    field(doc.add_paragraph(), 'TOC \\o "1-2" \\h \\z \\u', 'Update this field (right-click, Update Field) to show the table of contents.')
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    lines = body.split('\n'); i = 0; buf = []

    def flush():
        if buf:
            bm.add_inline(doc.add_paragraph(), ' '.join(s.strip() for s in buf)); buf.clear()

    while i < len(lines):
        ln = lines[i]
        if ln.startswith('```'):
            flush(); j = i + 1
            while not lines[j].startswith('```'):
                j += 1
            p = doc.add_paragraph(); r = p.add_run('\n'.join(lines[i + 1:j])); r.font.name = 'Consolas'; r.font.size = Pt(8.5)
            i = j + 1; continue
        m = re.match(r'^(#{1,4}) (.*)$', ln)
        if m:
            flush(); level = len(m.group(1))
            if level == 1 and m.group(2).startswith(BREAK_BEFORE + ('1. Introduction',)) and m.group(2) != 'Abstract':
                doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            doc.add_heading(m.group(2), level=min(level, 3)); i += 1; continue
        if ln.startswith('$$'):
            flush(); latex = ln.strip()
            while not (latex.endswith('$$') and len(latex) > 2):
                i += 1; latex += ' ' + lines[i].strip()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                node, tag = bm.omml(latex[2:-2]); p._p.append(node)
                if tag:
                    p.add_run(f'\t({tag})')
            except Exception:
                p.add_run(latex[2:-2]).italic = True
            i += 1; continue
        m = re.match(r'^!\[(.*?)\]\((.*?)\)(\{width=([\d.]+)in\})?', ln)
        if m:
            flush(); path = (HERE / m.group(2)).resolve()
            doc.add_picture(str(path), width=Inches(min(6.5, float(m.group(4) or 6))))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            if m.group(1):
                c = doc.add_paragraph(style='Caption'); bm.add_inline(c, m.group(1))
            i += 1; continue
        if ln.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s:|-]+\|$', lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            n = max(len(r) for r in rows)
            tb = doc.add_table(rows=len(rows), cols=n); tb.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci in range(n):
                    cell = tb.cell(ri, ci); cell.text = ''
                    bm.add_inline(cell.paragraphs[0], row[ci] if ci < len(row) else '')
                    cell.paragraphs[0].paragraph_format.space_after = Pt(0)
                    for rr in cell.paragraphs[0].runs:
                        rr.font.size = Pt(8.5); rr.bold = rr.bold or ri == 0
                    if ri == 0:
                        sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), 'E8E8E8'); cell._tc.get_or_add_tcPr().append(sh)
            trp = tb.rows[0]._tr.get_or_add_trPr(); hdr = OxmlElement('w:tblHeader'); hdr.set(qn('w:val'), 'true'); trp.append(hdr)
            doc.add_paragraph()
            continue
        m = re.match(r'^(\s*)([-*]|\d+\.) (.*)$', ln)
        if m:
            flush()
            style = 'List Number' if m.group(2)[0].isdigit() else 'List Bullet'
            if len(m.group(1)) >= 2:
                style += ' 2'
            bm.add_inline(doc.add_paragraph(style=style), m.group(3))
            i += 1; continue
        if not ln.strip() or ln.strip() == '---':
            flush(); i += 1; continue
        if re.match(r'^\*\*Table \d+\.\*\*', ln):
            flush(); p = doc.add_paragraph(); p.paragraph_format.keep_with_next = True; bm.add_inline(p, ln); i += 1; continue
        buf.append(ln); i += 1
    flush()
    import datetime
    cp = doc.core_properties; cp.author = meta.get('author', ''); cp.last_modified_by = meta.get('author', '')
    cp.title = meta.get('title', '') + ': ' + meta.get('subtitle', ''); cp.comments = meta.get('version', '')
    cp.keywords = 'human-AI collaboration; learning to defer; Bayesian decision-making; POMDP; risk'
    cp.created = cp.modified = datetime.datetime(2026, 10, 29)
    doc.save(str(out))


def build_html(md: str, out: Path, bm):
    meta, body = _meta(md)
    toc, parts = [], []
    lines = body.split('\n'); i = 0; buf = []; lst = []

    def flush():
        if buf:
            parts.append('<p>' + bm.html_inline(' '.join(s.strip() for s in buf)) + '</p>'); buf.clear()

    def close_lists(depth=0):
        while len(lst) > depth:
            parts.append(f'</{lst.pop()}>')

    while i < len(lines):
        ln = lines[i]
        if not re.match(r'^(\s*)([-*]|\d+\.) ', ln) and not (ln.startswith('   ') and lst and ln.strip()):
            if ln.strip():
                close_lists()
        if ln.startswith('```'):
            flush(); close_lists(); j = i + 1
            while not lines[j].startswith('```'):
                j += 1
            parts.append('<pre>' + H.escape('\n'.join(lines[i + 1:j])) + '</pre>'); i = j + 1; continue
        m = re.match(r'^(#{1,4}) (.*)$', ln)
        if m:
            flush(); close_lists(); lv = len(m.group(1)); title = m.group(2); sid = _slug(title)
            cls = ' class="brk"' if lv == 1 and title.startswith(BREAK_BEFORE + ('1. Introduction',)) and title != 'Abstract' else ''
            parts.append(f'<h{lv + 1} id="{sid}"{cls}>{bm.html_inline(title)}</h{lv + 1}>')
            if lv <= 2:
                toc.append((lv, sid, title))
            i += 1; continue
        if ln.startswith('$$'):
            flush(); latex = ln.strip()
            while not (latex.endswith('$$') and len(latex) > 2):
                i += 1; latex += ' ' + lines[i].strip()
            tag = re.search(r'\\tag\{([^}]*)\}', latex)
            src = re.sub(r'\\tag\{[^}]*\}', '', latex[2:-2]).replace(r'\rm ', r'\mathrm ')
            mml = bm.latex2mathml.converter.convert(src, display='block')
            parts.append(f'<div class="eq"><span class="eqbody">{mml}</span><span class="tag">{"(" + tag.group(1) + ")" if tag else ""}</span></div>'); i += 1; continue
        m = re.match(r'^!\[(.*?)\]\((.*?)\)(\{width=([\d.]+)in\})?', ln)
        if m:
            flush(); close_lists(); src = (HERE / m.group(2)).resolve().as_uri()
            parts.append(f'<figure><img src="{src}" style="width:{m.group(4) or 6}in;max-width:100%"><figcaption>{bm.html_inline(m.group(1))}</figcaption></figure>')
            i += 1; continue
        if ln.startswith('|'):
            flush(); close_lists(); rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s:|-]+\|$', lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            t = ['<table><thead><tr>' + ''.join(f'<th>{bm.html_inline(c)}</th>' for c in rows[0]) + '</tr></thead><tbody>']
            t += ['<tr>' + ''.join(f'<td>{bm.html_inline(c)}</td>' for c in r) + '</tr>' for r in rows[1:]]
            parts.append(''.join(t) + '</tbody></table>'); continue
        m = re.match(r'^(\s*)([-*]|\d+\.) (.*)$', ln)
        if m:
            flush(); depth = 1 + len(m.group(1)) // 2; kind = 'ol' if m.group(2)[0].isdigit() else 'ul'
            close_lists(depth)
            if len(lst) < depth:
                parts.append(f'<{kind}>'); lst.append(kind)
            parts.append(f'<li>{bm.html_inline(m.group(3))}</li>'); i += 1; continue
        if not ln.strip() or ln.strip() == '---':
            flush(); i += 1; continue
        if re.match(r'^\*\*Table \d+\.\*\*', ln):
            flush(); parts.append(f'<p class="tcap">{bm.html_inline(ln)}</p>'); i += 1; continue
        buf.append(ln); i += 1
    flush(); close_lists()
    toc_html = '<nav class="toc"><h2>Contents</h2><ul>' + ''.join(
        f'<li class="l{lv}"><a href="#{sid}">{bm.html_inline(t)}</a></li>' for lv, sid, t in toc) + '</ul></nav>'
    title = (f'<section class="titlepage"><h1 class="title">{H.escape(meta.get("title", ""))}</h1>'
             + ''.join(f'<p class="meta {k}">{H.escape(meta[k])}</p>' for k in ('subtitle', 'author', 'version', 'date') if k in meta) + '</section>')
    css = ('@page{size:A4;margin:22mm 20mm 20mm 20mm;@bottom-center{content:counter(page);font-family:"Times New Roman",serif;font-size:9pt;color:#444}}'
           '@page:first{@bottom-center{content:none}}'
           'body{font-family:"Times New Roman",serif;font-size:10.8pt;line-height:1.42;color:#111}'
           '.titlepage{page-break-after:always;text-align:center;padding-top:55mm}h1.title{font-size:30pt;margin-bottom:8mm}'
           'p.meta{margin:2mm 0}p.subtitle{font-style:italic;font-size:13.5pt;margin:0 15mm 8mm}p.author{font-size:13pt}p.version,p.date{font-size:10.5pt;color:#333}'
           '.toc{page-break-after:always}.toc ul{list-style:none;padding-left:0}.toc li.l1{margin-top:3px;font-weight:bold}.toc li.l2{margin-left:18px;font-size:9.8pt}'
           '.toc a{color:#111;text-decoration:none}a{color:#1a4f8a}'
           'h2{font-size:15pt;margin-top:1.3em;page-break-after:avoid}h2.brk{page-break-before:always}h3{font-size:12.3pt;page-break-after:avoid}h4,h5{font-size:11pt;page-break-after:avoid}'
           'table{border-collapse:collapse;font-size:8.4pt;margin:.3em 0 1em;width:100%;page-break-inside:auto}tr{page-break-inside:avoid}thead{display:table-header-group}'
           'th,td{border:1px solid #999;padding:2px 4px;vertical-align:top}td code{white-space:normal;overflow-wrap:anywhere;font-size:8pt}th{background:#ececec}'
           'p.tcap{margin:1em 0 .2em;page-break-after:avoid}'
           '.eq{display:flex;justify-content:center;align-items:center;margin:.6em 0;position:relative;page-break-inside:avoid}.eq .tag{position:absolute;right:0}'
           'pre{font-size:8.3pt;background:#f5f5f5;padding:6px;white-space:pre-wrap;page-break-inside:avoid}'
           'figure{text-align:center;margin:1em 0;page-break-inside:avoid}figcaption{font-style:italic;font-size:9.3pt;margin-top:.3em;text-align:left}code{font-size:9.3pt}'
           'li{margin:.15em 0}')
    out.write_text(f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{H.escape(meta.get("title", ""))}: final consolidated paper</title>'
                   f'<style>{css}</style></head><body>{title}{toc_html}' + '\n'.join(parts) + '</body></html>', encoding='utf-8')


def build_pdf(html_path: Path, pdf_path: Path):
    edge = r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe'
    subprocess.run([edge, '--headless=new', '--disable-gpu', '--no-pdf-header-footer', '--allow-file-access-from-files',
                    f'--print-to-pdf={pdf_path}', html_path.resolve().as_uri()], check=True, timeout=300)


def word_refresh(path: Path) -> bool:
    """Optional: let Microsoft Word populate the TOC and page fields, then clear the update-on-open flag."""
    import tempfile, zipfile, shutil
    ps = (f"$d='{path}'; $w=New-Object -ComObject Word.Application; $w.Visible=$false; $w.DisplayAlerts=0; "
          "try {{ $doc=$w.Documents.Open($d); $doc.TablesOfContents | ForEach-Object {{ $_.Update() }}; $doc.Fields.Update() | Out-Null; "
          "$doc.Save(); $doc.Close() }} finally {{ $w.Quit() }}").format()
    try:
        subprocess.run(['powershell', '-NoProfile', '-Command', ps], check=True, timeout=180)
    except Exception as e:
        print('Word refresh unavailable:', e); return False
    tmp = Path(tempfile.mkdtemp()) / path.name
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, 'w', zipfile.ZIP_DEFLATED) as zout:
        for it in zin.infolist():
            data = zin.read(it.filename)
            if it.filename == 'word/settings.xml':
                data = re.sub(rb'<w:updateFields[^>]*/>', b'', data)
            if it.filename == 'docProps/core.xml':   # Word stamps the local Office user; keep only the published author
                data = re.sub(rb'<cp:lastModifiedBy>.*?</cp:lastModifiedBy>', b'<cp:lastModifiedBy>Pratik653</cp:lastModifiedBy>', data)
            if it.filename == 'docProps/app.xml':
                data = re.sub(rb'<(Company|Manager)>.*?</>', b'', data)
            zout.writestr(it, data)
    shutil.copyfile(tmp, path); return True


if __name__ == '__main__':
    md = assemble()
    (HERE / f'{NAME}.md').write_text(md, encoding='utf-8', newline='\n')
    print('wrote', HERE / f'{NAME}.md', len(md), 'chars')
    if '--md-only' not in sys.argv:
        bm = _load_v2_helpers()
        build_docx(md, HERE / f'{NAME}.docx', bm)
        if '--word-refresh' in sys.argv:
            print('Word refresh:', word_refresh(HERE / f'{NAME}.docx'))
        tmp = HERE / '_build'; tmp.mkdir(exist_ok=True)
        build_html(md, tmp / 'paper.html', bm)
        build_pdf(tmp / 'paper.html', HERE / f'{NAME}.pdf')
        print('wrote DOCX and PDF')
