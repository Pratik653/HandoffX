"""Build the v0.2 manuscript from PART_II_v0.2.md + the unchanged v0.1 manuscript.

Outputs HandoffX_Updated_Research.md, .docx (native Word equations) and .pdf
(HTML with MathML printed by headless Microsoft Edge). Tooling only, not part of
the handoffx package. Requires python-docx, latex2mathml and lxml, Microsoft
Office's MML2OMML.XSL (LaTeX -> Word equations) and Microsoft Edge.

    python publication_v2/build_manuscript.py
"""
from __future__ import annotations
from pathlib import Path
import re
import latex2mathml.converter
from lxml import etree
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
XSL = Path(r'C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL')
_xslt = etree.XSLT(etree.parse(str(XSL)))


def assemble() -> str:
    """Front matter and Part II (v0.2) + the unchanged v0.1 manuscript as Part I."""
    part2 = (HERE / 'PART_II_v0.2.md').read_text(encoding='utf-8')
    fm, body2 = re.match(r'^---\n(.*?)\n---\n(.*)$', part2, re.S).groups()
    orig = (ROOT / '01_Publication/HandoffX_Research_Paper.md').read_text(encoding='utf-8')
    orig = re.sub(r'^---\n.*?\n---\n', '', orig, flags=re.S)
    for h in ('Abstract', 'Research integrity statement', 'Contents', 'Figures and tables', 'Notation and abbreviations'):
        orig = re.sub(rf'^# {h}$', f'# Part I. {h} (v0.1)', orig, count=1, flags=re.M)
    head, rest = body2.split('# 15. Motivation for v0.2', 1)
    guide = ('# How to read this document\n\nPart I is the complete v0.1 study, reproduced without change '
             '(Sections 1 to 14, its references and Appendices A to D). Part II (Sections 15 to 24) is the v0.2 '
             'extension: the diagnosis of the v0.1 result, the cross-episode controller, the pre-registered held-out '
             'benchmark, the agent demonstration and revised conclusions. Where Part I and Part II differ, Part II '
             'reports the later evidence; Part I conclusions about v0.1 remain valid for v0.1.\n\n'
             '# PART I. The v0.1 baseline study (unchanged)\n\n')
    return f'---\n{fm}\n---\n' + head + guide + orig + '\n\n# PART II. The v0.2 extension\n\n# 15. Motivation for v0.2' + rest


def omml(latex: str):
    tag = re.search(r'\\tag\{([^}]*)\}', latex)
    latex = re.sub(r'\\tag\{[^}]*\}', '', latex).strip()
    latex = latex.replace(r'\rm ', r'\mathrm ').replace(r'\,', ' ')
    mml = etree.fromstring(latex2mathml.converter.convert(latex).encode())
    return _xslt(mml).getroot(), (tag.group(1) if tag else None)


INLINE = re.compile(r'(\$[^$]+\$|\*\*[^*]+\*\*|\*[^*\s][^*]*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))')


def add_inline(par, text: str):
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('$') and tok.endswith('$') and len(tok) > 2:
            try:
                par._p.append(omml(tok[1:-1])[0])
            except Exception:
                par.add_run(tok[1:-1]).italic = True
        elif tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith('*') and tok.endswith('*'):
            par.add_run(tok[1:-1]).italic = True
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 1:
            r = par.add_run(tok[1:-1]); r.font.name = 'Consolas'
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)', tok):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok)
            par.add_run(m.group(1) if m.group(2).startswith(('#', '..', '.')) else f'{m.group(1)} ({m.group(2)})')
        else:
            par.add_run(tok.replace('\\_', '_').replace('\\*', '*'))


def build(md: str, out: Path):
    doc = Document()
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
    fm, body = re.match(r'^---\n(.*?)\n---\n(.*)$', md, re.S).groups()
    meta = dict(re.findall(r'^(\w+):\s*"?(.*?)"?$', fm, re.M))
    t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(meta.get('title', '')); r.bold = True; r.font.size = Pt(20)
    for k in ('subtitle', 'author', 'date'):
        if k in meta:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run(meta[k]).italic = k == 'subtitle'
    lines = body.split('\n'); i = 0; buf = []

    def flush():
        if buf:
            add_inline(doc.add_paragraph(), ' '.join(s.strip() for s in buf)); buf.clear()

    while i < len(lines):
        ln = lines[i]
        if ln.startswith('```'):
            flush(); j = i + 1
            while not lines[j].startswith('```'):
                j += 1
            p = doc.add_paragraph(); r = p.add_run('\n'.join(lines[i + 1:j])); r.font.name = 'Consolas'; r.font.size = Pt(9)
            i = j + 1; continue
        m = re.match(r'^(#{1,4}) (.*)$', ln)
        if m:
            flush(); level = len(m.group(1))
            if m.group(2).startswith('PART '):
                doc.add_page_break()
            doc.add_heading(m.group(2), level=min(level, 3)); i += 1; continue
        if ln.startswith('$$'):
            flush(); latex = ln.strip()
            while not (latex.endswith('$$') and len(latex) > 2):
                i += 1; latex += ' ' + lines[i].strip()
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            try:
                node, tag = omml(latex[2:-2])
                p._p.append(node)
                if tag:
                    p.add_run(f'    ({tag})')
            except Exception:
                p.add_run(latex[2:-2]).italic = True
            i += 1; continue
        m = re.match(r'^!\[(.*?)\]\((.*?)\)(\{width=([\d.]+)in\})?', ln)
        if m:
            flush(); path = (HERE / m.group(2)).resolve()
            doc.add_picture(str(path), width=Inches(float(m.group(4) or 6)))
            if m.group(1):
                c = doc.add_paragraph(); c.add_run(m.group(1)).italic = True
            i += 1; continue
        if ln.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s:|-]+\|$', lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            n = max(len(r) for r in rows)
            tb = doc.add_table(rows=len(rows), cols=n); tb.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci in range(n):
                    cell = tb.cell(ri, ci); cell.text = ''
                    add_inline(cell.paragraphs[0], row[ci] if ci < len(row) else '')
                    for rr in cell.paragraphs[0].runs:
                        rr.font.size = Pt(8.5); rr.bold = rr.bold or ri == 0
            continue
        m = re.match(r'^(\s*)([-*]|\d+\.) (.*)$', ln)
        if m:
            flush()
            add_inline(doc.add_paragraph(style='List Number' if m.group(2)[0].isdigit() else 'List Bullet'), m.group(3))
            i += 1; continue
        if not ln.strip() or ln.strip() == '---':
            flush(); i += 1; continue
        buf.append(ln); i += 1
    flush(); doc.save(str(out))


# ---------------------------------------------------------------- HTML / PDF
def html_inline(text: str) -> str:
    import html as h
    out = []
    for tok in INLINE.split(text):
        if not tok:
            continue
        if tok.startswith('$') and tok.endswith('$') and len(tok) > 2:
            out.append(latex2mathml.converter.convert(re.sub(r'\tag\{[^}]*\}', '', tok[1:-1]).replace(r'\rm ', r'\mathrm ')))
        elif tok.startswith('**') and tok.endswith('**') and len(tok) > 4:
            out.append(f'<strong>{h.escape(tok[2:-2])}</strong>')
        elif tok.startswith('*') and tok.endswith('*') and len(tok) > 2:
            out.append(f'<em>{h.escape(tok[1:-1])}</em>')
        elif tok.startswith('`') and tok.endswith('`') and len(tok) > 1:
            out.append(f'<code>{h.escape(tok[1:-1])}</code>')
        elif re.fullmatch(r'\[[^\]]+\]\([^)]+\)', tok):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', tok); out.append(f'<a href="{h.escape(m.group(2))}">{h.escape(m.group(1))}</a>')
        else:
            out.append(h.escape(tok.replace('\_', '_')))
    return ''.join(out)


def build_html(md: str, out: Path):
    import html as h
    fm, body = re.match(r'^---\n(.*?)\n---\n(.*)$', md, re.S).groups()
    meta = dict(re.findall(r'^(\w+):\s*"?(.*?)"?$', fm, re.M))
    parts = [f'<h1 class="title">{h.escape(meta.get("title", ""))}</h1>']
    for k in ('subtitle', 'author', 'date'):
        if k in meta:
            parts.append(f'<p class="meta {k}">{h.escape(meta[k])}</p>')
    lines = body.split('\n'); i = 0; buf = []; lst = None

    def flush():
        nonlocal lst
        if buf:
            parts.append('<p>' + html_inline(' '.join(s.strip() for s in buf)) + '</p>'); buf.clear()

    def close_list():
        nonlocal lst
        if lst:
            parts.append(f'</{lst}>'); lst = None

    while i < len(lines):
        ln = lines[i]
        if not re.match(r'^(\s*)([-*]|\d+\.) ', ln):
            close_list()
        if ln.startswith('```'):
            flush(); j = i + 1
            while not lines[j].startswith('```'):
                j += 1
            parts.append('<pre>' + h.escape('\n'.join(lines[i + 1:j])) + '</pre>'); i = j + 1; continue
        m = re.match(r'^(#{1,4}) (.*)$', ln)
        if m:
            flush(); lv = len(m.group(1)); cls = ' class="part"' if m.group(2).startswith('PART ') else ''
            parts.append(f'<h{lv + 1}{cls}>{html_inline(m.group(2))}</h{lv + 1}>'); i += 1; continue
        if ln.startswith('$$'):
            flush(); latex = ln.strip()
            while not (latex.endswith('$$') and len(latex) > 2):
                i += 1; latex += ' ' + lines[i].strip()
            tag = re.search(r'\tag\{([^}]*)\}', latex)
            mml = latex2mathml.converter.convert(re.sub(r'\tag\{[^}]*\}', '', latex[2:-2]).replace(r'\rm ', r'\mathrm '), display='block')
            parts.append(f'<div class="eq">{mml}<span class="tag">{"(" + tag.group(1) + ")" if tag else ""}</span></div>'); i += 1; continue
        m = re.match(r'^!\[(.*?)\]\((.*?)\)(\{width=([\d.]+)in\})?', ln)
        if m:
            flush(); src = (HERE / m.group(2)).resolve().as_uri()
            parts.append(f'<figure><img src="{src}" style="width:{m.group(4) or 6}in;max-width:100%"><figcaption>{html_inline(m.group(1))}</figcaption></figure>')
            i += 1; continue
        if ln.startswith('|'):
            flush(); rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[\s:|-]+\|$', lines[i]):
                    rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            t = ['<table>'] + ['<tr>' + ''.join(f'<{"th" if ri == 0 else "td"}>{html_inline(c)}</{"th" if ri == 0 else "td"}>' for c in r) + '</tr>'
                               for ri, r in enumerate(rows)] + ['</table>']
            parts.append(''.join(t)); continue
        m = re.match(r'^(\s*)([-*]|\d+\.) (.*)$', ln)
        if m:
            flush(); kind = 'ol' if m.group(2)[0].isdigit() else 'ul'
            if lst != kind:
                close_list(); parts.append(f'<{kind}>'); lst = kind
            parts.append(f'<li>{html_inline(m.group(3))}</li>'); i += 1; continue
        if not ln.strip() or ln.strip() == '---':
            flush(); i += 1; continue
        buf.append(ln); i += 1
    flush(); close_list()
    css = ('@page{size:A4;margin:22mm 20mm}body{font-family:"Times New Roman",serif;font-size:11pt;line-height:1.4}'
           'h1.title{text-align:center;font-size:22pt}p.meta{text-align:center}p.subtitle{font-style:italic}'
           'h2{font-size:15pt;margin-top:1.4em}h2.part{page-break-before:always;font-size:18pt}h3{font-size:12.5pt}h4,h5{font-size:11pt}'
           'table{border-collapse:collapse;font-size:8.5pt;margin:.6em 0;width:100%}th,td{border:1px solid #888;padding:2px 4px;vertical-align:top}'
           'th{background:#eee}.eq{display:flex;justify-content:center;align-items:center;gap:2em;margin:.6em 0}'
           'pre{font-size:8.5pt;background:#f6f6f6;padding:6px;white-space:pre-wrap}figure{text-align:center;margin:1em 0}'
           'figcaption{font-style:italic;font-size:9.5pt}code{font-size:9.5pt}')
    out.write_text(f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{h.escape(meta.get("title", ""))}</title>'
                   f'<style>{css}</style></head><body>' + '\n'.join(parts) + '</body></html>', encoding='utf-8')


def build_pdf(html_path: Path, pdf_path: Path):
    import subprocess
    edge = r'C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe'
    subprocess.run([edge, '--headless=new', '--disable-gpu', '--no-pdf-header-footer', '--allow-file-access-from-files',
                    f'--print-to-pdf={pdf_path}', html_path.resolve().as_uri()], check=True, timeout=300)


if __name__ == '__main__':
    md = assemble()
    (HERE / 'HandoffX_Updated_Research.md').write_text(md, encoding='utf-8', newline='\n')
    build(md, HERE / 'HandoffX_Updated_Research.docx')
    tmp = HERE / '_build'; tmp.mkdir(exist_ok=True)
    build_html(md, tmp / 'manuscript.html')
    build_pdf(tmp / 'manuscript.html', HERE / 'HandoffX_Updated_Research.pdf')
    print('wrote', HERE / 'HandoffX_Updated_Research.pdf')
